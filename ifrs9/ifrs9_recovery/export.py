"""Stage 7: registers, missingness, formula-based Excel master, summary and handoff archive."""
from __future__ import annotations

import shutil
import subprocess
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analysis, readiness, versions
from . import config as C
from . import eligibility as E
from .queue import split_urls
from .ratios import DEFINITIONS
from .util import number, read_csv, read_json, sha256_file, utcnow, write_csv, write_json

OUTCOMES = {
    "roa": "ROA on average assets (%)", "eqta": "Equity / assets (%)", "ocf_ratio": "OCF / average assets (%)",
    "allr": "Loan allowance / gross loans (%)", "leverage": "Liabilities / assets (%)",
    "car": "Reported CAR (%)", "car_computed": "Capital / RWA (%)", "npl_reported": "Reported NPL ratio (%)",
    "npl_calc": "NPL amount / gross loans (%)", "broad_coverage": "Broad NPL coverage (%)",
    "stage3_coverage": "Stage 3 coverage (%)", "impairment_assets": "Total impairment / average assets (%)",
    "liquidity_used": "Admitted domestic liquidity ratio (%)", "gcu": "Going-concern uncertainty code",
}
INPUT_FIELDS = ["pat", "assets", "liabilities", "equity", "ocf", "opening_assets", "gross_loans", "loan_allowance",
                "net_loans", "total_impairment", "loan_impairment", "npl_amount", "car", "regulatory_capital", "rwa",
                "car_minimum", "liquidity", "gcu", "stage3_gross", "stage3_allowance"]
FORMULA_INPUTS = ["bank", "year", "tier", "regime", "pat", "assets", "liabilities", "equity", "ocf", "opening_assets",
                  "gross_loans", "loan_allowance", "net_loans", "total_impairment", "npl_amount", "regulatory_capital",
                  "rwa", "stage3_gross", "stage3_allowance", "car", "car_minimum", "liquidity", "liquidity_admissible",
                  "gcu", "source_url", "version_note", "origin"]


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def missingness(panel: list[dict]) -> list[dict]:
    elig = E.summarise()
    out = []
    for f, label in OUTCOMES.items():
        avail = [r for r in panel if not _blank(r.get(f))]
        pre = sum(int(r["year"]) < C.ADOPTION_YEAR for r in avail)
        out.append({"outcome": f, "label": label, "definition": DEFINITIONS.get(f, ""),
                    "n_available": len(avail), "n_missing": len(panel) - len(avail), "n_rows": len(panel),
                    "pct_available": round(100 * len(avail) / len(panel), 1) if panel else 0,
                    "pre_2018_available": pre, "post_2018_available": len(avail) - pre,
                    "banks_with_any": len({r["bank"] for r in avail}),
                    "n_eligible_inferential": elig.get("liquidity" if f == "liquidity_used" else f, {}).get("eligible", "not an eligibility outcome")})
    return out


def field_coverage(before: list[dict], after: list[dict]) -> list[dict]:
    rows = []
    for f in INPUT_FIELDS:
        b = sum(not _blank(r.get(f)) for r in before)
        a = sum(not _blank(r.get(f)) for r in after)
        rows.append({"field": f, "nonblank_before": b, "rows_before": len(before),
                     "nonblank_after": a, "rows_after": len(after), "added": a - b})
    return rows


def unresolved(panel_keys: set, changes: list[dict], queue: list[dict], dl: dict, disc: dict) -> list[dict]:
    filled = {(c["bank"], str(c["year"]), c["field"]) for c in changes if c["kind"] in ("fill_blank", "correction")}
    job_by_bank_year: dict[tuple[str, str], list[str]] = {}
    job_by_url = {}
    for j in queue:
        for y in filter(None, j["target_years"].split(";")):
            job_by_bank_year.setdefault((j["bank"], y), []).append(j["job_id"])
        if j["url"]:
            job_by_url[j["url"]] = j["job_id"]

    def job_status(jid: str) -> str:
        if jid in dl:
            return f"{dl[jid]['status']}"
        st = disc.get("jobs", {}).get(jid)
        return st["status"] if st else "not_attempted"

    out = []
    for g in read_csv(C.TASKS / "bank_year_gaps.csv"):
        if (g["bank"], int(g["year"])) in panel_keys:
            continue
        jobs = job_by_bank_year.get((g["bank"], g["year"]), [])
        out.append({"gap_type": "bank_year", "bank": g["bank"], "year": g["year"], "field": "*core row*",
                    "priority": g["priority"], "status_before": g["current_status"],
                    "status_after": "not_admitted",
                    "retrieval_attempts": "; ".join(f"{j}:{job_status(j)}" for j in jobs) or "none",
                    "next_action": g["task"]})
    for f in read_csv(C.TASKS / "field_gaps_selected.csv"):
        if (f["bank"], f["year"], f["field"]) in filled:
            continue
        jids = [job_by_url[u] for u in split_urls(f["existing_source_url"]) if u in job_by_url]
        jid = next((j for j in jids if job_status(j) == "downloaded"), jids[0] if jids else "")
        out.append({"gap_type": f["task_kind"], "bank": f["bank"], "year": f["year"], "field": f["field"],
                    "priority": 1 if f["bank"] in ("Wema", "Union") else 6, "status_before": "blank in selected row",
                    "status_after": "unprocessed (source not retrieved)" if not jid or job_status(jid) != "downloaded"
                    else "source retrieved; awaiting page review",
                    "retrieval_attempts": "; ".join(f"{j}:{job_status(j)}" for j in jids) or "no existing source URL",
                    "next_action": f["instruction"]})
    return out


def source_manifest(queue, dl) -> list[dict]:
    rows: dict[str, dict] = {}
    for s in read_csv(C.SEEDS / "report_sources.csv"):
        rows.setdefault(s["url"], {"url": s["url"], "source_ids": s["source_id"], "bank": s["bank"],
                                   "lead_status_in_handoff": s["status"], "page_hint": s["page_hint"], "origin": "seed"})
    for s in read_csv(C.WORK / "update_pack" / "prior_evidence" / "Sources.csv"):
        rows.setdefault(s["url"], {"url": s["url"], "source_ids": s["source_id"], "bank": "Union",
                                   "lead_status_in_handoff": s["status"], "page_hint": "", "origin": "prior Sources.csv"})
    for r in read_csv(C.SELECTED_CSV):
        u = r.get("source_url", "").strip()
        if u:
            rows.setdefault(u, {"url": u, "source_ids": r.get("sources", ""), "bank": r["bank"],
                                "lead_status_in_handoff": "selected-row source", "page_hint": "", "origin": "selected CSV"})
    by_url = {v["requested_url"]: v for v in dl.values()}
    jobs = {j["url"]: j for j in queue if j["url"]}
    for u, r in rows.items():
        d = by_url.get(u, {})
        r.update({"queue_job": jobs.get(u, {}).get("job_id", ""), "retrieval_status": d.get("status", "not_attempted"),
                  "retrieval_reason": d.get("reason", ""), "final_url": d.get("final_url", ""),
                  "bytes": d.get("bytes", ""), "pdf_sha256": d.get("sha256", ""),
                  "local_pdf_path": d.get("local_pdf_path", ""), "accessed_utc": d.get("accessed_utc", "")})
    return list(rows.values())


# ---------------------------------------------------------------- Excel master

HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
NOTE = Font(italic=True, color="555555")


def _sheet(wb, title, rows, cols, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = NOTE
        r0 = 3
    for j, c in enumerate(cols, 1):
        cell = ws.cell(r0, j, c)
        cell.font, cell.fill = HDR, HDR_FILL
    for i, row in enumerate(rows, r0 + 1):
        for j, c in enumerate(cols, 1):
            v = row.get(c, "")
            if isinstance(v, str) and v.strip() and c not in ("bank", "year_text"):
                try:
                    n = number(v)
                    v = n if n is not None and c not in ("source_url", "url", "final_url", "evidence_id", "query") else v
                except ValueError:
                    pass
            ws.cell(i, j, None if v == "" else v)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = min(45, max(10, len(c) + 2))
    return ws, r0


def _ratio_formulas(col: dict[str, str], r: int) -> dict[str, str]:
    c = lambda f: f"Panel_Inputs!{col[f]}{r}"
    two = lambda n, d: f'=IF(AND(ISNUMBER({n}),ISNUMBER({d})),IF({d}>0,100*{n}/{d},""),"")'
    return {
        "average_assets": f'=IF(AND(ISNUMBER({c("assets")}),ISNUMBER({c("opening_assets")})),IF(AND({c("assets")}>0,{c("opening_assets")}>0),({c("assets")}+{c("opening_assets")})/2,""),"")',
        "roa": f'=IF(AND(ISNUMBER({c("pat")}),ISNUMBER(C{r})),IF(C{r}>0,100*{c("pat")}/C{r},""),"")',
        "eqta": two(c("equity"), c("assets")),
        "ocf_ratio": f'=IF(AND(ISNUMBER({c("ocf")}),ISNUMBER(C{r})),IF(C{r}>0,100*{c("ocf")}/C{r},""),"")',
        "allr": two(c("loan_allowance"), c("gross_loans")),
        "car_computed": two(c("regulatory_capital"), c("rwa")),
        "impairment_assets": f'=IF(AND(ISNUMBER({c("total_impairment")}),ISNUMBER(C{r})),IF(C{r}>0,100*{c("total_impairment")}/C{r},""),"")',
        "broad_coverage": two(c("loan_allowance"), c("npl_amount")),
        "npl_calc": two(c("npl_amount"), c("gross_loans")),
        "stage3_ratio": two(c("stage3_gross"), c("gross_loans")),
        "stage3_coverage": two(c("stage3_allowance"), c("stage3_gross")),
        "roa_closing": two(c("pat"), c("assets")),
        "leverage": two(c("liabilities"), c("assets")),
        "ln_assets": f'=IF(ISNUMBER({c("assets")}),IF({c("assets")}>0,LN({c("assets")}),""),"")',
        "balance_residual": f'=IF(AND(ISNUMBER({c("assets")}),ISNUMBER({c("liabilities")}),ISNUMBER({c("equity")})),{c("assets")}-{c("liabilities")}-{c("equity")},"")',
        "loan_residual": f'=IF(AND(ISNUMBER({c("gross_loans")}),ISNUMBER({c("loan_allowance")}),ISNUMBER({c("net_loans")})),{c("gross_loans")}-{c("loan_allowance")}-{c("net_loans")},"")',
        "liquidity_used": f'=IF(AND(ISNUMBER({c("liquidity")}),{c("liquidity_admissible")}=1),{c("liquidity")},"")',
        "car_buffer": f'=IF(AND(ISNUMBER({c("car")}),ISNUMBER({c("car_minimum")})),{c("car")}-{c("car_minimum")},"")',
    }


RATIO_ORDER = ["average_assets", "roa", "eqta", "ocf_ratio", "allr", "car_computed", "impairment_assets",
               "broad_coverage", "npl_calc", "stage3_ratio", "stage3_coverage", "roa_closing", "leverage",
               "ln_assets", "balance_residual", "loan_residual", "liquidity_used", "car_buffer"]


def build_workbook(path: Path, panel, sheets: dict[str, tuple[list[dict], list[str], str]], readme: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read_Me"
    for i, line in enumerate(readme, 1):
        for j, v in enumerate(line, 1):
            ws.cell(i, j, v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 120
    ws["A1"].font = Font(bold=True, size=14)

    pi, r0 = _sheet(wb, "Panel_Inputs", panel, FORMULA_INPUTS,
                    "Admitted operating-bank inputs. NGN million; percentage points. Blank = not admitted, never zero.")
    col = {f: get_column_letter(i) for i, f in enumerate(FORMULA_INPUTS, 1)}
    cr = wb.create_sheet("Calculated_Ratios")
    cr.cell(1, 1, "Formulas reference Panel_Inputs. Python-recomputed values are compared in Formula_Check.").font = NOTE
    head = ["bank", "year"] + RATIO_ORDER
    for j, h in enumerate(head, 1):
        cell = cr.cell(3, j, h)
        cell.font, cell.fill = HDR, HDR_FILL
    for i in range(len(panel)):
        r = r0 + 1 + i
        cr.cell(r, 1, f"=Panel_Inputs!A{r}")
        cr.cell(r, 2, f"=Panel_Inputs!B{r}")
        f = _ratio_formulas(col, r)
        for j, k in enumerate(RATIO_ORDER, 3):
            cr.cell(r, j, f[k])
    cr.freeze_panes = "C4"
    # Formula check sheet: python values vs Excel formulas.
    fc = wb.create_sheet("Formula_Check")
    fc.cell(1, 1, "Python value (from verified_panel.csv) and difference against the workbook formula.").font = NOTE
    chk = [k for k in RATIO_ORDER if k not in ("car_buffer",)]
    hdr = ["bank", "year"] + [f"{k}_py" for k in chk] + [f"{k}_diff" for k in chk]
    for j, h in enumerate(hdr, 1):
        cell = fc.cell(3, j, h)
        cell.font, cell.fill = HDR, HDR_FILL
    for i, row in enumerate(panel):
        r = r0 + 1 + i
        fc.cell(r, 1, row["bank"])
        fc.cell(r, 2, int(row["year"]))
        for j, k in enumerate(chk):
            v = number(row.get(k)) if not _blank(row.get(k)) else None
            pc = 3 + j
            fc.cell(r, pc, v)
            xl = f"Calculated_Ratios!{get_column_letter(3 + RATIO_ORDER.index(k))}{r}"
            py = f"{get_column_letter(pc)}{r}"
            fc.cell(r, 3 + len(chk) + j,
                    f'=IF(AND(ISNUMBER({xl}),ISNUMBER({py})),{xl}-{py},IF(AND({xl}="",ISBLANK({py})),0,"MISMATCH"))')
    for title, (rows, cols, note) in sheets.items():
        _sheet(wb, title, rows, cols, note)
    wb.save(path)


def recalc_with_libreoffice(path: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    out = path.parent / "_recalc"
    out.mkdir(exist_ok=True)
    try:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                       check=True, capture_output=True, timeout=240)
    except (subprocess.SubprocessError, OSError):
        return False
    produced = out / path.name
    if produced.is_file():
        shutil.move(str(produced), str(path))
        shutil.rmtree(out, ignore_errors=True)
        return True
    return False


def formula_audit(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb["Formula_Check"]
    diffs = errors = checked = 0
    hdr = [c.value for c in ws[3]]
    for row in ws.iter_rows(min_row=4, values_only=True):
        for h, v in zip(hdr, row):
            if h and str(h).endswith("_diff"):
                checked += 1
                if v == "MISMATCH" or (isinstance(v, str) and v.startswith("#")):
                    errors += 1
                elif isinstance(v, (int, float)) and abs(v) > 1e-6:
                    diffs += 1
    cr = wb["Calculated_Ratios"]
    xl_errors = sum(1 for row in cr.iter_rows(min_row=4, values_only=True) for v in row
                    if isinstance(v, str) and v.startswith("#"))
    return {"cells_checked": checked, "numeric_differences": diffs, "mismatches": errors, "excel_errors": xl_errors}


# ---------------------------------------------------------------- main

def run() -> dict:
    C.ensure_dirs()
    # One eligibility ruleset feeds every analytical output below (summary, workbook, analysis CSVs, readiness).
    versions.build()
    E.build()
    analysis_out = analysis.build()
    readiness.build()
    before = read_csv(C.SELECTED_CSV)
    panel = read_csv(C.DATA / "verified_panel.csv")
    changes = read_csv(C.DATA / "change_register.csv")
    queue = read_csv(C.DATA / "task_queue.csv")
    dl = read_json(C.STATE / "download_state.json", {}).get("reports", {})
    disc = read_json(C.STATE / "discover_state.json", {})
    vstate = read_json(C.DATA / "validate_state.json", {})
    inv = read_json(C.DATA / "inventory.json", {})
    budget = read_json(C.STATE / "nimble_budget.json", {"attempted_requests": 0})

    keys = {(r["bank"], int(r["year"])) for r in panel}
    miss = missingness(panel)
    cov = field_coverage(before, panel)
    gaps = unresolved(keys, changes, queue, dl, disc)
    manifest = source_manifest(queue, dl)
    write_csv(C.DATA / "missingness_by_outcome.csv", miss)
    write_csv(C.DATA / "field_coverage.csv", cov)
    write_csv(C.DATA / "unresolved_gaps.csv", gaps)
    write_csv(C.DATA / "source_manifest.csv", manifest)

    # Indexed report archive (only real downloaded PDFs).
    archive = [{"bank": v["bank"], "job_id": k, "source_url": v["final_url"], "sha256": v["sha256"], "bytes": v["bytes"],
                "local_pdf_path": v["local_pdf_path"], "accessed_utc": v["accessed_utc"]}
               for k, v in dl.items() if v.get("status") == "downloaded"]
    write_csv(C.REPORTS / "index.csv", archive,
              ["bank", "job_id", "source_url", "sha256", "bytes", "local_pdf_path", "accessed_utc"])

    # Copy the rendered pages that back admitted values into a committed evidence folder.
    ev_dir = C.OUTPUTS / "evidence_pages"
    ev_dir.mkdir(parents=True, exist_ok=True)
    adm_path = C.DATA / "review" / "admissions.csv"
    copied = set()
    for r in (read_csv(adm_path) if adm_path.is_file() else []):
        src = C.ROOT / r.get("rendered_page_path", "")
        if r.get("rendered_page_path") and src.is_file():
            dst = ev_dir / ("_".join(Path(r["rendered_page_path"]).parts[-2:]))
            if dst.name not in copied:
                shutil.copy2(src, dst)
                copied.add(dst.name)

    from . import summary as S
    n_base = len(before)
    new_rows = vstate.get("new_bank_years", [])
    corr = [c for c in changes if c["kind"] == "correction"]
    fills = [c for c in changes if c["kind"] == "fill_blank"]
    recon = read_csv(C.DATA / "reconciliation_checks.csv")
    exc = read_csv(C.DATA / "exceptions_register.csv")
    conflicts = read_csv(C.DATA / "conflicts.csv")
    backing = read_csv(C.DATA / "evidence_backing.csv")
    evidence = read_csv(C.DATA / "evidence_long.csv")
    gap_bank_years = [g for g in gaps if g["gap_type"] == "bank_year"]
    gap_fields = [g for g in gaps if g["gap_type"] != "bank_year"]
    stats = S.release_stats()
    pst, rcs, gcs = stats["panel"], stats["reconciliation"], stats["going_concern"]
    progress = [
        {"measure": "Selected bank-years", "value": pst["rows"]},
        {"measure": "  baseline retained / additional", "value": f"{pst['baseline_rows_retained']} / {pst['new_rows']}"},
        {"measure": "  pre-adoption / post-adoption", "value": f"{pst['pre_adoption']} / {pst['post_adoption']}"},
        {"measure": "Banks represented", "value": pst["banks"]},
        {"measure": "Blank baseline cells filled (logged)", "value": stats["changes"].get("fill_blank", 0)},
        {"measure": "Corrections to non-blank baseline values", "value": stats["changes"].get("correction", 0)},
        {"measure": "Admissions accepted / rejected", "value": f"{stats['admissions']['accepted']} / {stats['admissions']['rejected']}"},
        {"measure": "Reconciliation pass / fail / documented exception", "value": f"{rcs.get('pass', 0)} / {rcs.get('fail', 0)} / {rcs.get('documented_source_exception', 0)}"},
        {"measure": "Evidence lineage matched / checked", "value": f"{stats['evidence_lineage']['matched']} / {stats['evidence_lineage']['checked']}"},
        {"measure": "Going-concern reviews recorded (coded 1)", "value": f"{gcs['reviews_recorded']} ({gcs['reviews_coded_1']})"},
        {"measure": "Panel GCU coded / unknown", "value": f"{gcs['panel_coded']} / {gcs['panel_unknown']}"},
        {"measure": "Remaining bank-year positions", "value": stats["remaining"]["bank_years"]},
        {"measure": "Remaining field requests in selected rows", "value": stats["remaining"]["field_requests"]},
        {"measure": "PDFs archived and hashed (not in zip)", "value": len(archive)},
        {"measure": "Nimble REST attempts / plugin invocations", "value": f"{stats['nimble'].get('rest_client', {}).get('attempted_requests', 0)} / {stats['nimble'].get('mcp_plugin', {}).get('invocations', 0)}"},
    ]
    write_csv(C.OUTPUTS / "progress_table.csv", progress, ["measure", "value"])

    readme = [
        ["IFRS 9 VERIFIED MASTER"],
        ["Generated (UTC)", utcnow()],
        ["Baseline", f"{n_base} selected bank-years (inputs/IFRS9_Selected_Data_Updated_20261010.csv, unchanged; SHA-256 verified)."],
        ["This run", f"{len(new_rows)} new bank-years admitted; {len(fills)} blank fields filled; {len(corr)} corrections. See Progress."],
        ["Units", "Monetary amounts NGN million; ratios in percentage points. Blank = not admitted or unavailable, never zero."],
        ["Entity rule", "Historical Nigerian operating bank only; never holding company, consolidated group, foreign subsidiary or segment."],
        ["Formulas", "Calculated_Ratios holds live formulas over Panel_Inputs; Formula_Check compares them with the Python recomputation."],
        ["Evidence", "Evidence_Long: inherited records are labelled inherited_not_rechecked_this_run; only reviewed admissions are verified_reported."],
        ["Caution", "A report link is not a verified value; a verified value is not a complete bank-year; no regressions are estimated here."],
    ]
    for d in sorted(DEFINITIONS):
        readme.append([d, DEFINITIONS[d]])

    sheets = {
        "Progress": (progress, ["measure", "value"], "Measured progress for this run."),
        "Field_Coverage": (cov, list(cov[0].keys()), "Non-blank input cells before (baseline CSV) and after."),
        "Missingness": (miss, list(miss[0].keys()), "Available observations by outcome; n_eligible_inferential from the eligibility ruleset."),
        "Change_Register": (changes, ["bank", "year", "field", "old", "new", "kind", "basis", "reason"], "Before/after log."),
        "Reconciliation": (recon, ["bank", "year", "check", "residual", "tolerance", "rounding_unit_ngnm", "passed", "status", "exception_detail", "note"],
                           "Status: pass / fail / documented_source_exception. Tolerances from printed precision and input rounding."),
        "Field_Status": (read_csv(C.DATA / "field_status.csv"), list(read_csv(C.DATA / "field_status.csv")[0].keys()),
                         "Basis of every non-blank cell: baseline_inherited, reported_text_layer, manual_transcription, derived_from_verified, restated comparative, recomputed_ratio."),
        "Version_Decisions": (read_csv(C.DATA / "review" / "version_decisions.csv"), list(read_csv(C.DATA / "review" / "version_decisions.csv")[0].keys()),
                              "Primary rule: original-year report; alternatives classified and kept for sensitivity."),
        "Source_Alternatives": (read_csv(C.DATA / "review" / "source_alternatives.csv"), ["evidence_id", "bank_key", "financial_year", "variable", "raw_value", "normalised_value_ngnm_or_pct", "pdf_page_1_based", "printed_page", "verification_status", "review_reason"],
                                "Printed values not selected (conflicting subtotals, original presentations, later comparatives)."),
        "Going_Concern": (read_csv(C.DATA / "review" / "going_concern_reviews.csv"), list(read_csv(C.DATA / "review" / "going_concern_reviews.csv")[0].keys()),
                          "Reviews of original-year auditor reports; blank code = not reviewed (unknown, not zero)."),
        "Document_Status": (read_csv(C.DATA / "review" / "document_status.csv"), list(read_csv(C.DATA / "review" / "document_status.csv")[0].keys()),
                            "Document type and audit-report coverage per downloaded file."),
        "Denominators": (read_csv(C.DATA / "denominator_compatibility.csv") or [{"bank": ""}], list((read_csv(C.DATA / "denominator_compatibility.csv") or [{"bank": ""}])[0].keys()),
                         "Opening assets vs prior-year closing assets in the panel."),
        "Outcome_Eligibility": (read_csv(E.PATH), E.COLUMNS,
                                "Single eligibility ruleset (ifrs9_recovery.eligibility): descriptive / longitudinal / inferential use per bank-year and outcome, with rule ids."),
        "Outcome_Summary": (read_csv(C.OUTPUTS / "analysis" / "outcome_summary.csv"), list(read_csv(C.OUTPUTS / "analysis" / "outcome_summary.csv")[0].keys()),
                            "Eligible N, banks, pre/post, Tier 1/2, exclusions per analysis dataset (outputs/analysis/)."),
        "Row_Versions": (read_csv(C.DATA / "row_versions.csv"), list(read_csv(C.DATA / "row_versions.csv")[0].keys()),
                         "Primary reporting version per bank-year; fallback rows are primary but dropped in sensitivity."),
        "Exceptions": (exc, ["id", "bank", "year", "field", "expected", "actual", "status", "rule"], "Known exceptions."),
        "Conflicts": (conflicts, ["type", "bank", "year", "field", "value_a", "value_b", "detail"], "Unresolved conflicts."),
        "Unresolved_Gaps": (gaps, list(gaps[0].keys()) if gaps else ["gap_type"], "Outstanding bank-years and field requests."),
        "Source_Register": (manifest, list(manifest[0].keys()), "Every known source URL and its retrieval outcome in this run."),
        "Evidence_Long": (evidence, list(evidence[0].keys()) if evidence else ["evidence_id"], "Long-form evidence."),
        "Evidence_Backing": (backing, list(backing[0].keys()) if backing else ["bank"], "Does each selected value match an evidence record?"),
        "Task_Queue": (queue, list(queue[0].keys()), "Executable queue in priority order."),
        "Candidates": (read_csv(C.DATA / "candidates.csv"), ["bank", "year", "field", "value", "candidate_source", "locator", "status"],
                       "Unverified raw candidates. Not used in Panel_Inputs."),
    }
    xlsx = C.OUTPUTS / "IFRS9_Verified_Master.xlsx"
    build_workbook(xlsx, panel, sheets, readme)
    recalculated = recalc_with_libreoffice(xlsx)
    audit = formula_audit(xlsx) if recalculated else {"note": "LibreOffice unavailable; formulas not cached"}

    write_json(C.OUTPUTS / "workbook_formula_audit.json", {"recalculated": recalculated, **audit})
    stats = S.release_stats()
    write_json(C.OUTPUTS / "release_stats.json", stats)
    write_json(C.OUTPUTS / "release_manifest.json", S.release_manifest(stats))
    summary = C.OUTPUTS / "reconciliation_summary.md"
    # Two passes: the first creates every file the zip will hold, so the counts in the final summary are final.
    for _ in range(2):
        handoff = build_handoff_zip()
        summary.write_text(S.render(stats, audit, handoff), encoding="utf-8")
    handoff = build_handoff_zip()
    return {"panel_rows": len(panel), "workbook": str(xlsx.relative_to(C.ROOT)), "recalculated": recalculated,
            "formula_audit": audit, "handoff": handoff,
            "remaining_bank_years": len(gap_bank_years), "remaining_field_requests": len(gap_fields)}


def build_handoff_zip() -> dict:
    z = C.OUTPUTS / "research_data_handoff.zip"
    include = []
    for d in ("data", "outputs", "ifrs9_recovery", "tests", "tasks", "seeds", "checks"):
        for p in sorted((C.ROOT / d).rglob("*")):
            if p.is_file() and p != z and "__pycache__" not in p.parts and p.suffix != ".tmp" and "_recalc" not in p.parts:
                include.append(p)
    for f in ("requirements.txt", "requirements.lock", "README_PIPELINE.md", "CLAUDE_PROMPT.md", "START_HERE.md",
              ".env.example", ".gitignore", "reports/index.csv", "pytest.ini"):
        if (C.ROOT / f).is_file():
            include.append(C.ROOT / f)
    pdfs = read_csv(C.REPORTS / "index.csv") if (C.REPORTS / "index.csv").is_file() else []
    adm_path = C.DATA / "review" / "admissions.csv"
    ev_dir = C.OUTPUTS / "evidence_pages"
    needed = {"_".join(Path(r["rendered_page_path"]).parts[-2:]) for r in (read_csv(adm_path) if adm_path.is_file() else [])
              if r.get("rendered_page_path")}
    present = {p.name for p in ev_dir.glob("*.png")} if ev_dir.is_dir() else set()
    missing = sorted(needed - present)
    manifest = ["research_data_handoff.zip manifest", f"created_utc: {utcnow()}", "",
                "CONTENTS: code, tests, data registers, outputs, and rendered page images (outputs/evidence_pages/) for every admitted value.",
                f"Rendered evidence images included: {len(needed & present)}; admission image paths without an included image: {len(missing)}.",
                "Image name = <first 12 hex of PDF SHA-256>_p<PDF page>.png, rendered at 110 dpi with PyMuPDF.",
                f"SOURCE PDFs ARE NOT INCLUDED ({len(pdfs)} files). Retrieve each from the URL in reports/index.csv and verify its SHA-256.",
                "Manual and image-only transcriptions cannot be reproduced from a hash alone: re-inspect the cited PDF page.",
                "Original inputs are identified by checks/snapshot_manifest.json.", "", "files:"]
    manifest += [f"  MISSING IMAGE: {m}" for m in missing]
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in include:
            rel = p.relative_to(C.ROOT).as_posix()
            zf.write(p, rel)
            manifest.append(f"  {rel}  sha256={sha256_file(p)}")
        zf.writestr("MANIFEST.txt", "\n".join(manifest) + "\n")
    return {"files": len(include) + 1, "evidence_images": len(needed & present), "missing_images": len(missing), "pdfs": len(pdfs)}
