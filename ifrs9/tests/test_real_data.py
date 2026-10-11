"""Real-data regression tests: the saved release outputs and the actual edge cases found during review.

These read the committed data registers (not synthetic fixtures). Tests that need the downloaded
source PDFs are skipped, and reported as not run, when those files are absent from the checkout.
"""
import csv
import json
import re
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
REVIEW = DATA / "review"
OUT = ROOT / "outputs"


def rows(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def num(x):
    try:
        return float(str(x).replace(",", ""))
    except ValueError:
        return None


@pytest.fixture(scope="module")
def panel():
    return {(r["bank"], int(r["year"])): r for r in rows(DATA / "verified_panel.csv")}


@pytest.fixture(scope="module")
def admissions():
    return rows(REVIEW / "admissions.csv")


@pytest.fixture(scope="module")
def checks():
    return {(r["bank"], int(r["year"]), r["check"]): r for r in rows(DATA / "reconciliation_checks.csv")}


@pytest.fixture(scope="module")
def stats():
    from ifrs9_recovery import summary
    return summary.release_stats()


# 1. Release snapshot: the narrative is generated from the saved CSVs and agrees with them.

def test_summary_counts_reconcile_with_csvs(panel, stats):
    changes = rows(DATA / "change_register.csv")
    pre = sum(1 for r in panel.values() if r["regime"] == "IAS 39")
    assert stats["panel"]["rows"] == len(panel)
    assert stats["panel"]["pre_adoption"] == pre and stats["panel"]["post_adoption"] == len(panel) - pre
    assert stats["changes"]["new_bank_year"] == sum(c["kind"] == "new_bank_year" for c in changes)
    assert stats["changes"]["fill_blank"] == sum(c["kind"] == "fill_blank" for c in changes)
    assert set(stats["changes"]) <= {"new_bank_year", "fill_blank", "correction"}
    coded = [r for r in panel.values() if r["gcu"].strip() != ""]
    assert stats["going_concern"]["panel_coded"] == len(coded)
    assert stats["going_concern"]["panel_unknown"] == len(panel) - len(coded)
    recon = rows(DATA / "reconciliation_checks.csv")
    assert stats["reconciliation"]["run"] == len(recon)
    assert stats["reconciliation"].get("fail", 0) == sum(r["status"] == "fail" for r in recon)
    assert stats["reconciliation"]["pass"] == sum(r["status"] == "pass" for r in recon)


def test_summary_text_does_not_contradict_going_concern_fills(stats):
    text = (OUT / "reconciliation_summary.md").read_text(encoding="utf-8")
    fills = sum(c["kind"] == "fill_blank" and c["field"] == "gcu" for c in rows(DATA / "change_register.csv"))
    assert fills > 0
    assert "No going-concern code was added" not in text
    assert "no original-year auditor report was reviewed" not in text.lower()
    assert f"{fills} filled in baseline rows" in text
    reviews = sum(r["report_reviewed"] not in ("not reviewed", "not coded") and not r["report_reviewed"].startswith("partially")
                  for r in rows(REVIEW / "going_concern_reviews.csv"))
    assert f"auditor reports reviewed and recorded: {reviews}" in text


def test_every_changed_baseline_cell_is_in_the_change_register(panel):
    baseline = rows(ROOT / "inputs" / "IFRS9_Selected_Data_Updated_20261010.csv")
    logged = {(c["bank"], int(c["year"]), c["field"]) for c in rows(DATA / "change_register.csv")}
    unlogged = []
    for b in baseline:
        key = (b["bank"], int(b["year"]))
        p = panel[key]
        for f, old in b.items():
            new = p.get(f, "")
            if old == new or (num(old) is not None and num(new) is not None and abs(num(old) - num(new)) < 1e-9):
                continue
            if (*key, f) not in logged:
                unlogged.append((*key, f, old, new))
    assert unlogged == []


def test_all_baseline_rows_retained_and_no_value_cleared(panel):
    for b in rows(ROOT / "inputs" / "IFRS9_Selected_Data_Updated_20261010.csv"):
        p = panel[(b["bank"], int(b["year"]))]
        cleared = [f for f, v in b.items() if v.strip() and not p.get(f, "").strip()]
        assert cleared == [], (b["bank"], b["year"], cleared)


def test_lineage_every_nonblank_input_matched_or_qualified():
    backing = rows(DATA / "evidence_backing.csv")
    assert backing and all(r["value_matches_evidence"] == "True" for r in backing)
    aliased = [r for r in backing if r["match_basis"].startswith("evidence recorded as")]
    assert aliased and all(r["qualification"].strip() for r in aliased)
    # The four alias matches investigated in this run are qualified, not given invented citations.
    named = {(r["bank"], r["year"], r["field"]) for r in backing if r["qualification"].strip()}
    for k in [("First Bank", "2013", "car"), ("First Bank", "2014", "car"), ("Unity", "2018", "car"), ("Unity", "2019", "car")]:
        assert k in named


# 3. Individual reconciliation exceptions.

def test_uba_2018_car_passes_on_display_precision_only(panel, checks):
    c = checks[("UBA", 2018, "reported CAR vs capital/RWA")]
    assert c["status"] == "pass"
    assert num(c["tolerance"]) < 0.51  # half a whole-percent step plus input rounding, not a global tolerance
    assert "printed literal '20'" in c["note"]
    p = panel[("UBA", 2018)]
    assert num(p["car"]) == 20 and num(p["car_transitional"]) == 22.08
    assert num(p["regulatory_capital"]) == 245482 and num(p["rwa"]) == 1248509
    assert abs(100 * 245482 / 1248509 - 19.66201285) < 1e-6


def test_wema_2021_car_stays_a_documented_source_exception(panel, checks):
    c = checks[("Wema", 2021, "reported CAR vs capital/RWA")]
    assert c["status"] == "documented_source_exception" and c["passed"] == "False"
    assert num(c["tolerance"]) <= 0.0051  # the UBA precision rule does not relax a 2-dp printed ratio
    assert num(panel[("Wema", 2021)]["car"]) == 11.71
    stats_eligible = json.loads((OUT / "release_stats.json").read_text())["endpoints"]["car"]
    assert stats_eligible["excluded_from_inference"] >= 1


def test_stanbic_2020_capital_is_derived_and_printed_subtotal_preserved(panel, admissions):
    rc = next(r for r in admissions if r["evidence_id"] == "STAN20-RC")
    assert rc["value_basis"] == "derived_from_verified" and rc["raw_value"] == "derived"
    assert rc["derive_formula"] == "sum" and set(rc["component_ids"].split(";")) == {"STAN20-T1", "STAN20-T2"}
    assert num(panel[("Stanbic", 2020)]["regulatory_capital"]) == 221408
    alts = {r["evidence_id"]: r for r in rows(REVIEW / "source_alternatives.csv")}
    assert alts["ALT-STAN20-RC-PRINTED"]["raw_value"] == "176,542"
    assert alts["ALT-STAN20-RCT-PRINTED"]["raw_value"] == "182,720"
    dec = [d for d in rows(REVIEW / "version_decisions.csv")
           if (d["bank"], d["year"]) == ("Stanbic", "2020") and d["field"] in ("regulatory_capital", "capital_transitional")]
    assert dec and all("unresolved" in d["decision"] for d in dec)
    fs = {(r["bank"], r["year"], r["field"]): r for r in rows(DATA / "field_status.csv")}
    assert fs[("Stanbic", "2020", "regulatory_capital")]["value_basis"].startswith("derived_from_verified")


# 4. First Bank: numbers supported, audit coverage unverified, gcu unknown.

def test_first_bank_audit_coverage(panel):
    # FY2017 extract has no auditor's report: 2017 stays unknown. 2016 coded only from the full original 2016 report.
    doc = next(d for d in rows(REVIEW / "document_status.csv") if d["job_id"] == "xs-809800e65f")
    assert doc["document_type"] == "issuer-hosted financial-statement extract"
    assert doc["audit_report_reviewed_for_gcu"] == "unverified"
    assert panel[("First Bank", 2017)]["gcu"].strip() == ""
    r = next(r for r in rows(REVIEW / "going_concern_reviews.csv") if (r["bank"], r["year"]) == ("First Bank", "2016"))
    assert r["job_id"] == "xs-bbae41452e" and r["gcu_code"] == "0"
    assert panel[("First Bank", 2016)]["gcu"].strip() in ("0", "0.0")
    for y in (2016, 2017):
        assert num(panel[("First Bank", y)]["assets"]) > 0


# 6. Version policy.

def test_wema_2022_impairment_retained_with_reclassification_decision(panel):
    assert num(panel[("Wema", 2022)]["total_impairment"]) == 4762.274
    d = next(d for d in rows(REVIEW / "version_decisions.csv")
             if (d["bank"], d["year"], d["field"]) == ("Wema", "2022", "total_impairment"))
    assert d["classification"] == "labelled reclassification" and "4808.387" in d["alternative_value"]


def test_every_version_decision_is_classified():
    allowed = re.compile(r"^(labelled restatement|labelled reclassification|unexplained difference|source internal inconsistency"
                         r"|possible scope incompatibility.*|confirmed scope difference.*|corroborated by original-year.*|fallback: .+)$")
    for d in rows(REVIEW / "version_decisions.csv"):
        assert allowed.match(d["classification"]), d


def test_denominator_pairs_documented_not_smoothed():
    for d in rows(DATA / "denominator_compatibility.csv"):
        if d["compatible"] != "True":
            assert d["opening_version"] and d["prior_closing_version"]


# 7. Going concern.

def test_unity_2019_material_uncertainty_confirmed_against_note(panel):
    r = next(r for r in rows(REVIEW / "going_concern_reviews.csv") if (r["bank"], r["year"]) == ("Unity", "2019"))
    assert r["gcu_code"] == "1" and r["mu_section"].lower().startswith("yes")
    assert "35" in r["management_gc_disclosure"]
    assert panel[("Unity", 2019)]["gcu"].strip() in ("1", "1.0")


def test_gc_codes_only_from_reviews_covering_the_bank(panel):
    reviews = {(r["bank"], int(r["year"])): r for r in rows(REVIEW / "going_concern_reviews.csv")}
    for c in rows(DATA / "change_register.csv"):
        if c["field"] == "gcu":
            rv = reviews[(c["bank"], int(c["year"]))]
            from ifrs9_recovery.validate import GC_APPLICABLE
            assert rv["report_reviewed"] != "not reviewed" and rv["opinion_applicability"] in GC_APPLICABLE
            assert rv["auditor"].strip() and rv["report_date"].strip() and rv["signature_page_pdf"].strip()


def test_not_reviewed_is_unknown_not_zero(panel):
    for r in rows(REVIEW / "going_concern_reviews.csv"):
        if r["report_reviewed"] == "not reviewed":
            assert r["gcu_code"] == ""


# 5. Admission boundary on real admissions.

def test_sign_transforms_only_for_expense_variables_with_reason(admissions):
    from ifrs9_recovery.validate import SIGN_CONVERTIBLE
    for r in admissions:
        if r["sign_transform"] not in ("", "none"):
            assert SIGN_CONVERTIBLE.match(r["variable"]) and r["sign_reason"].strip(), r["evidence_id"]


def test_negative_pat_equity_ocf_are_preserved(admissions, panel):
    for r in admissions:
        if r["variable"] in ("pat", "equity", "ocf") and r["raw_value"] not in ("derived", ""):
            raw_neg = r["raw_value"].strip().startswith(("(", "-"))
            assert raw_neg == (num(r["normalised_value_ngnm_or_pct"]) < 0), r["evidence_id"]
    assert num(panel[("Unity", 2019)]["equity"]) < 0


def test_derived_admissions_recompute_from_components(admissions):
    from ifrs9_recovery.validate import check_derived
    by_id = {r["evidence_id"]: r for r in admissions if r.get("admit", "").lower() in ("yes", "true", "1")}
    derived = [r for r in by_id.values() if r["value_basis"] == "derived_from_verified"]
    assert derived
    for r in derived:
        assert check_derived(r, by_id) == [], r["evidence_id"]


def test_structured_scope_on_every_admission(admissions):
    from ifrs9_recovery.validate import check_scope
    for r in admissions:
        assert check_scope(r) == [], r["evidence_id"]


def _pdfs_present(admissions):
    return all((ROOT / r["local_pdf_path"]).is_file() for r in admissions)


def test_all_admissions_hash_page_and_image_verified(admissions):
    if not _pdfs_present(admissions):
        pytest.skip("source PDFs not in this checkout (retrieve via reports/index.csv)")
    from ifrs9_recovery import config as C
    from ifrs9_recovery.util import read_json
    from ifrs9_recovery.validate import check_admission
    dl = read_json(C.STATE / "download_state.json", {}).get("reports", {})
    downloads = {v["sha256"]: v for v in dl.values() if v.get("status") == "downloaded"}
    cohort = {r["bank"]: r for r in rows(C.TASKS / "cohort.csv")}
    bad = {r["evidence_id"]: errs for r in admissions if (errs := check_admission(r, downloads, cohort))}
    assert bad == {}


def test_handoff_includes_every_evidence_image():
    import zipfile
    z = OUT / "research_data_handoff.zip"
    names = set(zipfile.ZipFile(z).namelist())
    manifest = zipfile.ZipFile(z).read("MANIFEST.txt").decode()
    assert "SOURCE PDFs ARE NOT INCLUDED" in manifest and "MISSING IMAGE" not in manifest
    assert not any(n.endswith(".pdf") for n in names)


def test_no_secret_in_tracked_outputs():
    pat = re.compile(r"(?i)nimble[_-]?api[_-]?key\s*[=:]\s*['\"]?[A-Za-z0-9_\-]{16,}")
    for d in (DATA, OUT):
        for p in d.rglob("*"):
            if p.suffix in (".csv", ".json", ".md", ".txt"):
                assert not pat.search(p.read_text(encoding="utf-8", errors="ignore")), p


def test_fidelity_2016_car_documented_exception_not_resolved(panel, checks):
    c = checks[("Fidelity", 2016, "reported CAR vs capital/RWA")]
    assert c["status"] == "documented_source_exception" and c["passed"] == "False"
    assert num(panel[("Fidelity", 2016)]["car"]) == 17.22
    assert abs(100 * 196913 / 1138258 - 17.30) < 0.005


def test_stanbic_2018_writeback_stored_as_negative_charge(admissions, panel):
    r = next(a for a in admissions if a["evidence_id"] == "STAN18-TI")
    assert r["raw_value"] == "2,988" and num(r["normalised_value_ngnm_or_pct"]) == -2988
    assert r["sign_transform"] == "expense_presented_negative" and "write-back" in r["sign_reason"]


def test_npl_only_from_explicit_labels(admissions):
    # 'Individually impaired' (IAS 39) amounts were not admitted as NPL for the 2014-2020 Fidelity additions.
    assert not [a for a in admissions if a["evidence_id"] in ("FID14-NPL", "FID15-NPL", "FID16-NPL", "FID17-NPL", "FID20-NPL")]


def test_wema_2013_from_exchange_filing_audit_unverified_gcu_unknown(panel):
    assert num(panel[("Wema", 2013)]["assets"]) == 330872.475
    assert panel[("Wema", 2013)]["gcu"].strip() == ""
    doc = next(d for d in rows(REVIEW / "document_status.csv") if d["job_id"] == "xs-24238b1f7e")
    assert doc["audit_report_in_file"] == "No" and doc["audit_report_reviewed_for_gcu"] == "unverified"


def test_zenith_2013_group_only_filing_not_used(panel, admissions):
    # The NGX filing holds only consolidated group figures; Zenith 2013 comes from the full report's Bank columns.
    assert not [a for a in admissions if a["source_id"] == "xs-86dd272867"]
    z = [a for a in admissions if a["bank_key"] == "Zenith" and a["financial_year"] == "2013" and a["variable"] != "gcu"]
    assert z and all(a["source_id"] == "xs-44bcf809bd" and a["column_scope"] == "bank_separate" for a in z)


def test_unity_2021_material_uncertainty_coded_one(panel):
    r = next(r for r in rows(REVIEW / "going_concern_reviews.csv") if (r["bank"], r["year"]) == ("Unity", "2021"))
    assert r["gcu_code"] == "1" and r["mu_section"].lower().startswith(("yes", "true"))
    assert panel[("Unity", 2021)]["gcu"].strip() in ("1", "1.0")
    assert num(panel[("Unity", 2021)]["equity"]) < 0


def test_access_2020_full_impact_and_transitional_car_kept_separate(panel):
    p = panel[("Access", 2020)]
    assert num(p["car"]) == 15.87 and num(p["car_transitional"]) == 17.22


def test_unity_2016_capital_kam_flagged_not_coded(panel):
    r = next(r for r in rows(REVIEW / "going_concern_reviews.csv") if (r["bank"], r["year"]) == ("Unity", "2016"))
    assert r["gcu_code"] == "0" and "FLAG for sensitivity" in r["coding_basis"]


# ---- Phase 1/5/7: one eligibility ruleset enforced everywhere ----

ANALYSIS = OUT / "analysis"


@pytest.fixture(scope="module")
def elig():
    return rows(DATA / "outcome_eligibility.csv")


def _flag(r, k):
    return str(r[k]) in ("True", "true", "1")


def _excluded(elig_rows, outcome):
    return {(r["bank"], r["year"]) for r in elig_rows if r["outcome"] == outcome and _flag(r, "available") and not _flag(r, "inferential_use")}


def _workbook_sheet(name):
    from openpyxl import load_workbook
    ws = load_workbook(OUT / "IFRS9_Verified_Master.xlsx", read_only=True, data_only=True)[name]
    data = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, row in enumerate(data) if row and row[0] == "bank" or (row and row[0] == "dataset"))
    hdr = [str(h) for h in data[hdr_i]]
    return [dict(zip(hdr, [("" if v is None else v) for v in row])) for row in data[hdr_i + 1:] if row and row[0] not in (None, "")]


def test_union_2018_excluded_from_pooled_allowance_inference(elig):
    for o in ("allr", "broad_coverage"):
        r = next(r for r in elig if (r["bank"], r["year"], r["outcome"]) == ("Union", "2018", o))
        assert _flag(r, "available") and _flag(r, "descriptive_use")
        assert not _flag(r, "inferential_use") and not _flag(r, "longitudinal_use")
    a = next(r for r in rows(ANALYSIS / "allowance_gross_loans.csv") if (r["bank"], r["year"]) == ("Union", "2018"))
    assert a["inferential_use"] == "False" and "EL-UNION18-ALLOWANCE" in a["rule_ids"]


def test_wema_2021_2022_impairment_comparison_excluded(elig):
    w21 = next(r for r in elig if (r["bank"], r["year"], r["outcome"]) == ("Wema", "2021", "impairment_assets"))
    w22 = next(r for r in elig if (r["bank"], r["year"], r["outcome"]) == ("Wema", "2022", "impairment_assets"))
    assert not _flag(w21, "inferential_use") and not _flag(w21, "longitudinal_use")
    assert not _flag(w22, "change_comparable")
    d = next(d for d in rows(REVIEW / "version_decisions.csv") if (d["bank"], d["year"], d["field"]) == ("Wema", "2021", "total_impairment"))
    assert "108,256" in d["evidence"] and "1996.137" in d["sensitivity_alternative"]


def test_car_source_exceptions_not_inferential_and_not_reconciled(elig, checks):
    for b, y in (("Wema", 2021), ("Fidelity", 2016)):
        assert checks[(b, y, "reported CAR vs capital/RWA")]["status"] == "documented_source_exception"
        for o in ("car", "car_computed"):
            r = next(r for r in elig if (r["bank"], r["year"], r["outcome"]) == (b, str(y), o))
            assert not _flag(r, "inferential_use"), (b, y, o)


def test_stanbic_2020_derived_capital_never_labelled_reported_total(admissions):
    rc = next(a for a in admissions if a["evidence_id"] == "STAN20-RC")
    assert rc["verification_status"] == "derived_from_verified" and rc["raw_value"] == "derived"
    ev = [e for e in rows(DATA / "evidence_long.csv") if e["bank_key"] == "Stanbic" and e["financial_year"] == "2020" and e["variable"] == "regulatory_capital"]
    sel = [e for e in ev if e["evidence_id"] == "STAN20-RC"]
    assert sel and all(e["verification_status"] == "derived_from_verified" for e in sel)
    alt = [e for e in ev if e["evidence_id"] == "ALT-STAN20-RC-PRINTED"]
    assert not alt or all(e["verification_status"] != "verified_reported" or "176,542" in e["raw_value"] for e in alt)
    fs = next(f for f in rows(DATA / "field_status.csv") if (f["bank"], f["year"], f["field"]) == ("Stanbic", "2020", "regulatory_capital"))
    assert fs["value_basis"].startswith("derived_from_verified")


def test_full_impact_and_transitional_car_never_overwrite(admissions, panel):
    by = {}
    for a in admissions:
        by.setdefault((a["bank_key"], int(a["financial_year"])), {})[a["variable"]] = a
    n = 0
    for k, v in by.items():
        if "car" in v and "car_transitional" in v:
            n += 1
            assert v["car"]["evidence_id"] != v["car_transitional"]["evidence_id"]
            assert not v["car"]["evidence_id"].endswith("CART")
            assert num(panel[k]["car"]) == round(num(v["car"]["normalised_value_ngnm_or_pct"]), 6)
            assert num(panel[k]["car_transitional"]) == round(num(v["car_transitional"]["normalised_value_ngnm_or_pct"]), 6)
    assert n >= 3
    assert num(panel[("UBA", 2018)]["car"]) == 20 and num(panel[("UBA", 2018)]["car_transitional"]) == 22.08


def test_unknown_going_concern_never_zero(panel, elig):
    reviewed = {(r["bank"], r["year"]) for r in rows(REVIEW / "going_concern_reviews.csv") if r["gcu_code"] != ""}
    baseline = {(r["bank"], r["year"]) for r in rows(ROOT / "inputs" / "IFRS9_Selected_Data_Updated_20261010.csv") if r["gcu"].strip()}
    for (b, y), p in panel.items():
        if (b, str(y)) not in reviewed | baseline:
            assert p["gcu"].strip() == "", (b, y)
    gc = [r for r in elig if r["outcome"] == "gcu"]
    assert all(_flag(r, "available") == (r["value"] != "") for r in gc)
    assert len(rows(ANALYSIS / "going_concern.csv")) == sum(r["value"] != "" for r in gc)


def test_fallback_rows_keep_fallback_flags(elig):
    expected = {("Access", "2018"), ("Unity", "2013"), ("Unity", "2020"), ("Wema", "2021"), ("Union", "2015"), ("Fidelity", "2021"), ("GTBank", "2017")}
    rv = {(r["bank"], r["year"]): r for r in rows(DATA / "row_versions.csv")}
    assert expected <= {k for k, r in rv.items() if r["fallback_class"]}
    for k in expected:
        assert rv[k]["original_report_unavailable"] == "yes" and "drop in sensitivity" in rv[k]["treatment"]
        eq = next(r for r in elig if (r["bank"], r["year"], r["outcome"]) == (*k, "eqta"))
        assert _flag(eq, "sensitivity_drop") and eq["fallback_class"].startswith("fallback")
        a = next(r for r in rows(ANALYSIS / "equity_assets.csv") if (r["bank"], r["year"]) == k)
        assert a["sensitivity_drop"] == "True"


def test_outcome_summaries_match_exported_csv_row_counts(elig):
    import json
    from ifrs9_recovery.analysis import DATASETS
    stats = json.loads((OUT / "release_stats.json").read_text())["endpoints"]
    summ = {r["dataset"]: r for r in rows(ANALYSIS / "outcome_summary.csv")}
    wb = {r["dataset"]: r for r in _workbook_sheet("Outcome_Summary")}
    sizes = rows(OUT / "readiness" / "model_sample_sizes.csv")
    for name, (outcome, _) in DATASETS.items():
        csv_rows = rows(ANALYSIS / f"{name}.csv")
        inf = [r for r in csv_rows if r["inferential_use"] == "True"]
        assert int(summ[name]["descriptive_n"]) == len(csv_rows)
        assert int(summ[name]["eligible_n"]) == len(inf) == stats[outcome]["eligible"]
        assert int(wb[name]["eligible_n"]) == len(inf)
        pooled = next(s for s in sizes if s["outcome"] == outcome and s["model"].startswith("pooled"))
        assert int(pooled["N"]) == len(inf)


def test_same_rows_excluded_everywhere(elig):
    import json
    from ifrs9_recovery.analysis import DATASETS
    stats = json.loads((OUT / "release_stats.json").read_text())["endpoints"]
    wb = _workbook_sheet("Outcome_Eligibility")
    text = (OUT / "reconciliation_summary.md").read_text(encoding="utf-8")
    for name, (outcome, _) in DATASETS.items():
        e = _excluded(elig, outcome)
        a = {(r["bank"], r["year"]) for r in rows(ANALYSIS / f"{name}.csv") if r["inferential_use"] == "False"}
        s = {tuple(x.split(" (")[0].rsplit(" ", 1)) for x in stats[outcome]["excluded_rows"]}
        w = {(r["bank"], str(r["year"])) for r in wb if r["outcome"] == outcome and str(r["available"]) == "True" and str(r["inferential_use"]) == "False"}
        assert e == a == s == w, outcome
        for b, y in e:
            assert f"{b} {y}" in text


def test_access_2022_from_bank_columns_only():
    panel = [r for r in rows(DATA / "verified_panel.csv") if r["bank"] == "Access" and r["year"] == "2022"]
    assert len(panel) == 1
    r = panel[0]
    assert (num(r["assets"]), num(r["liabilities"]), num(r["equity"]), num(r["pat"])) == (12535280.0, 11466613.0, 1068667.0, 166660.0)
    assert num(r["assets"]) == num(r["liabilities"]) + num(r["equity"])
    assert r["gcu"] == ""  # auditor's report incomplete in the available PDF
    specs = [s for s in rows(REVIEW / "specs.csv") if s["bank"] == "Access" and s["year"] == "2022"]
    assert specs and all(s["column_scope"] == "bank_separate" and "Bank" in s["column_scope_basis"] and "Company" not in s["column_scope_basis"] for s in specs)


def test_holdco_sources_rejected_and_unused():
    scr = rows(REVIEW / "source_scope_screening.csv")
    rejected = {s["job_id"] for s in scr if s["decision"] == "rejected" and s["job_id"]}  # whole-document rejections
    assert rejected  # FBN Holdings 2018-2022 (subgroup column) and others
    used = {s["job_id"] for s in rows(REVIEW / "specs.csv")}
    assert not rejected & used
    fbn = [s for s in scr if s["source_id"].startswith("FBNH") and s["source_id"].endswith("_AR")]
    assert len(fbn) == 5 and all(s["decision"] == "rejected" and s["nigeria_operating_bank_only"] == "no" for s in fbn)


def test_holdco_disclosure_rows_are_labelled_and_reconcile():
    panel = rows(DATA / "verified_panel.csv")
    held = [r for r in panel if "HoldCo package" in r["scope"]]
    assert sorted((r["bank"], r["year"]) for r in held) == [("GTBank", "2021"), ("GTBank", "2022"), ("Stanbic", "2013"),
                                                          ("Stanbic", "2014"), ("Stanbic", "2015"), ("Stanbic", "2016"), ("Stanbic", "2017")]
    for r in held:
        assert abs(num(r["assets"]) - num(r["liabilities"]) - num(r["equity"])) < 0.01
        assert r["gcu"] == ""  # no auditor opinion on the operating bank
    specs = [s for s in rows(REVIEW / "specs.csv") if s["column_scope"] == "bank_subsidiary_disclosure"]
    assert specs and all(s["document_scope"] == "holding_company_report" for s in specs)


def test_new_full_report_rows_balance():
    panel = {(r["bank"], r["year"]): r for r in rows(DATA / "verified_panel.csv")}
    for y in (2013, 2014, 2015, 2018, 2019, 2020, 2021, 2022):
        r = panel[("Ecobank Nigeria", str(y))]
        assert abs(num(r["assets"]) - num(r["liabilities"]) - num(r["equity"])) < 0.01 and r["gcu"] == "0.0"
    for y in (2015, 2016):
        assert ("GTBank", str(y)) in panel


def test_first_bank_2018_2019_2022_from_bank_columns():
    panel = {(r["bank"], r["year"]): r for r in rows(DATA / "verified_panel.csv")}
    expect = {"2018": (3985907.0, 3543223.0, 442684.0, 56952.0, "0.0"), "2019": (4668279.0, 4128814.0, 539465.0, 43917.0, "0.0"),
              "2022": (8125060.0, 7372174.0, 752886.0, 91415.0, "")}
    for y, (a, l, e, pat, gcu) in expect.items():
        r = panel[("First Bank", y)]
        assert (num(r["assets"]), num(r["liabilities"]), num(r["equity"]), num(r["pat"])) == (a, l, e, pat)
        assert r["gcu"] == gcu  # 2022: only an ISA 810 summary-FS report, so not coded
    assert ("First Bank", "2020") not in panel and ("First Bank", "2021") not in panel
    specs = [s for s in rows(REVIEW / "specs.csv") if s["bank"] == "First Bank" and s["year"] in ("2018", "2019", "2022")]
    assert specs and all(s["column_scope"] == "bank_separate" and "BANK" in s["column_scope_basis"] for s in specs)
    assert all("comparative" in s["original_or_comparative"] for s in specs if s["year"] == "2022")
